#!/usr/bin/env python3
"""Convert HTML to DOCX using python-docx and BeautifulSoup"""
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys
from pathlib import Path

def html_to_docx(html_path, docx_path):
    with open(html_path, 'r', encoding='utf-8') as f:
        html_content = f.read()
    
    soup = BeautifulSoup(html_content, 'html.parser')
    
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    style.font.name = 'Lora'
    style.font.size = Pt(11)
    
    # Process elements
    for element in soup.body.find_all(True):
        if element.name == 'h1':
            doc.add_heading(element.get_text(), level=0)
        elif element.name == 'h2':
            doc.add_heading(element.get_text(), level=1)
        elif element.name == 'h3':
            doc.add_heading(element.get_text(), level=2)
        elif element.name == 'p':
            text = element.get_text().strip()
            if text:
                p = doc.add_paragraph(text)
                p.paragraph_format.line_spacing = 1.5
        elif element.name == 'pre' or element.name == 'code':
            text = element.get_text().strip()
            if text:
                p = doc.add_paragraph(text)
                p.style = 'Code'
        elif element.name in ['ul', 'ol']:
            for li in element.find_all('li'):
                text = li.get_text().strip()
                if text:
                    doc.add_paragraph(text, style='List Bullet' if element.name == 'ul' else 'List Number')
    
    # Save document
    doc.save(docx_path)
    print(f"DOCX generated: {docx_path}")
    print(f"DOCX size: {Path(docx_path).stat().st_size / 1024:.1f} KB")
    return docx_path

if __name__ == '__main__':
    html_path = sys.argv[1] if len(sys.argv) > 1 else 'out/_build/book.html'
    docx_path = sys.argv[2] if len(sys.argv) > 2 else 'out/Prompt-Video-IPA.docx'
    html_to_docx(html_path, docx_path)
